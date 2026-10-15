"""
Excel Template Generator for B.Sc. Nursing Examination Marksheets
Generates formatted .xlsx templates compatible with the dynamic ingestion parser.
"""

import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from backend.schemas.course_schemas import SEMESTER_DEFINITIONS


def generate_semester_template(semester_number: int = 1) -> bytes:
    """
    Generates a professionally styled Excel template matching SGPGIMS Lucknow
    Data Sheet format for the given semester.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Semester_{semester_number}_Data"

    sem_def = SEMESTER_DEFINITIONS.get(semester_number, SEMESTER_DEFINITIONS[1])

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Deep blue
    sub_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")     # Vibrant blue
    light_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")   # Alice blue

    title_font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="334155")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Row 1: College Header
    ws.cell(row=1, column=1, value="Sanjay Gandhi Postgraduate Institute of Medical Sciences, Lucknow")
    ws.cell(row=1, column=1).font = title_font

    # Row 2: Exam Session Title
    ws.cell(
        row=2,
        column=1,
        value=f"Data sheet {sem_def.degree_title} {sem_def.year_name} {sem_def.semester_name} Examination held in January - 2026"
    )
    ws.cell(row=2, column=1).font = subtitle_font

    # Rows 3 & 4: Table Headers
    # Demographic columns
    demographics_headers = [
        ("S.No.", 1),
        ("Name", 2),
        ("Father's Name", 3),
        ("Enroll No.", 4),
        ("Roll NO.", 5),
    ]

    for label, col in demographics_headers:
        cell = ws.cell(row=4, column=col, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    current_col = 6
    if semester_number == 1:
        # Course 1: Applied Anatomy & Applied Physiology
        ws.merge_cells(start_row=3, start_column=6, end_row=3, end_column=12)
        c1_title = ws.cell(row=3, column=6, value="Applied Anatomy & Applied Physiology (Code No. ANAT 105 + PHYS 110)")
        c1_title.font = header_font
        c1_title.fill = header_fill
        c1_title.alignment = Alignment(horizontal="center", vertical="center")

        c1_cols = [
            ("Int. 25", 6),
            ("Applied Anatomy Ex. 37", 7),
            ("Applied Physiology Ex. 38", 8),
            ("Total (75)", 9),
            ("Total (100)", 10),
            ("Letter Grade", 11),
            ("Grade Point", 12),
        ]
        for label, c in c1_cols:
            cell = ws.cell(row=4, column=c, value=label)
            cell.font = header_font
            cell.fill = sub_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Course 2: Applied Sociology & Applied Psychology
        ws.merge_cells(start_row=3, start_column=13, end_row=3, end_column=19)
        c2_title = ws.cell(row=3, column=13, value="Applied Sociology & Applied Psychology (Code No. SOCI 115 + PSYC 120)")
        c2_title.font = header_font
        c2_title.fill = header_fill
        c2_title.alignment = Alignment(horizontal="center", vertical="center")

        c2_cols = [
            ("Int. 25", 13),
            ("Applied Sociology Ex. 37", 14),
            ("Applied Psychology Ex. 38", 15),
            ("Total (75)", 16),
            ("Total (100)", 17),
            ("Letter Grade", 18),
            ("Grade Point", 19),
        ]
        for label, c in c2_cols:
            cell = ws.cell(row=4, column=c, value=label)
            cell.font = header_font
            cell.fill = sub_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Course 3: Communicative English
        ws.merge_cells(start_row=3, start_column=20, end_row=3, end_column=24)
        c3_title = ws.cell(row=3, column=20, value="Communicative English (Code No. ENGL 101)")
        c3_title.font = header_font
        c3_title.fill = header_fill
        c3_title.alignment = Alignment(horizontal="center", vertical="center")

        c3_cols = [
            ("Int. 25", 20),
            ("Ex. 25", 21),
            ("Total (50)", 22),
            ("Letter Grade", 23),
            ("Grade Point", 24),
        ]
        for label, c in c3_cols:
            cell = ws.cell(row=4, column=c, value=label)
            cell.font = header_font
            cell.fill = sub_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # Summary columns
        ws.cell(row=4, column=25, value="Grand Total (250)").font = header_font
        ws.cell(row=4, column=25).fill = header_fill
        ws.cell(row=4, column=25).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=4, column=25).border = thin_border

        ws.cell(row=4, column=26, value="Result / Remark").font = header_font
        ws.cell(row=4, column=26).fill = header_fill
        ws.cell(row=4, column=26).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.cell(row=4, column=26).border = thin_border

        # Sample student rows for Semester 1 (mirroring real SGPGIMS datasheet)
        sample_rows = [
            [
                1, "ADNAN RAFIQ KUMAR", "MOHMMAD RAFIQ KUMAR", "2025631", "202501003",
                19, 29, 25, 54, 73, "B+", 7,
                21, 35, 23, 58, 79, "A", 8,
                23, 21, 44, "O", 10,
                196, "PASS"
            ],
            [
                2, "AKASH VERMA", "MAHESH CHANDRA", "2025632", "202501004",
                19, 23, 22, 45, 64, "B", 6,
                18, 32, 26, 58, 76, "A", 8,
                22, 23, 45, "O", 10,
                185, "PASS"
            ],
            [
                3, "AMAN CHAUDHARY", "SURESH CHANDRA", "2025633", "202501005",
                15, 21, 20, 41, 56, "C", 5,
                15, 24, 13, "37+0.5=37.5", 53, "C", 5,
                20, 18, 38, "A", 8,
                147, "PASS WITH 0.5 GRACE MARKS IN APPLIED SOCIOLOGY & APPLIED PSYCHOLOGY"
            ],
            [
                4, "AMRITA PATEL", "DEVI DAS", "2025635", "202501007",
                21, 29, 33, 62, 83, "A+", 9,
                22, 26, 27, 53, 75, "A", 8,
                22, 23, 45, "O", 10,
                203, "PASS"
            ],
            [
                5, "VANSH KUMAR", "MUKESH KUMAR", "2025692", "202501064",
                13, 4, 5, 9, 22, "F", 0,
                15, 10, 9, 19, 34, "F", 0,
                13, 16, 29, "C", 5,
                85, "FAIL in Applied Anatomy & Applied Physiology and Applied Sociology & Applied Psychology"
            ]
        ]

        for r_idx, r_data in enumerate(sample_rows, start=5):
            for c_idx, val in enumerate(r_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                if c_idx in [1, 4, 5, 11, 12, 18, 19, 23, 24]:
                    cell.alignment = Alignment(horizontal="center")
                elif c_idx in [2, 3, 26]:
                    cell.alignment = Alignment(horizontal="left")
                else:
                    cell.alignment = Alignment(horizontal="right")

    else:
        # General semester layout generator
        c = 6
        for course in sem_def.courses:
            for comp in course.components:
                # Group header
                start_c = c
                ws.cell(row=3, column=start_c, value=f"{course.course_code}: {course.title} ({comp.name})").font = header_font
                ws.cell(row=3, column=start_c).fill = header_fill

                ws.cell(row=4, column=c, value=f"IA ({int(comp.ia_max)})").font = header_font
                ws.cell(row=4, column=c).fill = sub_fill
                ws.cell(row=4, column=c).border = thin_border
                c += 1

                ws.cell(row=4, column=c, value=f"EA ({int(comp.ea_max)})").font = header_font
                ws.cell(row=4, column=c).fill = sub_fill
                ws.cell(row=4, column=c).border = thin_border
                c += 1

                ws.cell(row=4, column=c, value=f"Total ({int(comp.total_max)})").font = header_font
                ws.cell(row=4, column=c).fill = sub_fill
                ws.cell(row=4, column=c).border = thin_border
                c += 1

                end_c = c - 1
                if end_c > start_c:
                    ws.merge_cells(start_row=3, start_column=start_c, end_row=3, end_column=end_c)

        # Summary header
        ws.cell(row=4, column=c, value=f"Grand Total ({int(sem_def.grand_total_marks)})").font = header_font
        ws.cell(row=4, column=c).fill = header_fill
        ws.cell(row=4, column=c).border = thin_border
        c += 1

        ws.cell(row=4, column=c, value="Result / Remark").font = header_font
        ws.cell(row=4, column=c).fill = header_fill
        ws.cell(row=4, column=c).border = thin_border

        # Sample row
        sample_student = [1, "PRIYANKA KUMARI", "RAMESH KUMAR", "2026101", "202601001"]
        for course in sem_def.courses:
            for comp in course.components:
                sample_student.extend([int(comp.ia_max * 0.8), int(comp.ea_max * 0.75), int(comp.total_max * 0.76)])
        sample_student.extend([int(sem_def.grand_total_marks * 0.76), "PASS"])

        for c_idx, val in enumerate(sample_student, start=1):
            cell = ws.cell(row=5, column=c_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
