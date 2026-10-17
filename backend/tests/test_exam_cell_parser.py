"""
Unit and Integration Tests for Examination Cell Ingestion Engine (Phase 1B)
"""

import io
import openpyxl
import pytest
from backend.parsers.exam_cell_parser import (
    ExamCellParserAdapter,
    unmerge_and_replicate_cells,
    parse_grace_expression,
    clean_demographic_id,
    clean_text_string,
    is_footer_marker,
)
from backend.schemas.canonical_models import StudentRecord


def test_grace_expression_extractor():
    # Case 1: Hyphen format "35+2.5-37.5"
    raw, grace, tot, ab = parse_grace_expression("35+2.5-37.5")
    assert raw == 35.0
    assert grace == 2.5
    assert tot == 37.5
    assert ab is False

    # Case 2: Equals format "37+0.5=37.5"
    raw, grace, tot, ab = parse_grace_expression("37+0.5=37.5")
    assert raw == 37.0
    assert grace == 0.5
    assert tot == 37.5
    assert ab is False

    # Case 3: Short addition "35+2.5"
    raw, grace, tot, ab = parse_grace_expression("35+2.5")
    assert raw == 35.0
    assert grace == 2.5
    assert tot == 37.5
    assert ab is False

    # Case 4: Absent indicators
    raw, grace, tot, ab = parse_grace_expression("Ab")
    assert ab is True
    assert raw is None

    raw2, grace2, tot2, ab2 = parse_grace_expression("ABSENT")
    assert ab2 is True

    # Case 5: Standard float / int / whitespace
    raw, grace, tot, ab = parse_grace_expression(" 20.5 ")
    assert raw == 20.5
    assert grace == 0.0
    assert tot == 20.5
    assert ab is False

    # Case 6: Hyphen / empty
    raw, grace, tot, ab = parse_grace_expression("-")
    assert raw is None
    assert grace == 0.0


def test_demographic_id_sanitization():
    # Float integer representation
    assert clean_demographic_id(202501001.0) == "202501001"
    assert clean_demographic_id("202501001.0") == "202501001"
    # Pure integer
    assert clean_demographic_id(2025630) == "2025630"
    # Scientific notation
    assert clean_demographic_id("2.02501002E+08") == "202501002"
    # Whitespace stripping
    assert clean_demographic_id("  202501002  ") == "202501002"
    # Null / None
    assert clean_demographic_id(None) == ""


def test_footer_marker_detection():
    assert is_footer_marker(["Prepared By", "Mr Vinod Kumar Gupta"]) is True
    assert is_footer_marker(["Checked by", "Dr. Radha K."]) is True
    assert is_footer_marker(["DEAN", "College of Nursing"]) is True
    assert is_footer_marker(["EXECUTIVE REGISTRAR"]) is True
    assert is_footer_marker(["Student Name", "ADITI SHARMA", "PASS"]) is False


def test_parser_with_grace_string_in_cell():
    """
    Constructs a synthetic workbook with raw grace string expressions
    e.g. '35+2.5-37.5' in cell, floating headers on row 3-4, and footer markers.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ExamTabulation"

    # Row 1-2: Institutional noise / titles
    ws.cell(row=1, column=1, value="Sanjay Gandhi Postgraduate Institute of Medical Sciences")
    ws.cell(row=2, column=1, value="Examination Cell Tabulation Sheet 2026")

    # Row 3: Parent Course Headers (Anchor detection) with merged cell spans
    ws.merge_cells(start_row=3, start_column=6, end_row=3, end_column=11)
    ws.cell(row=3, column=6, value="Applied Anatomy & Applied Physiology (ANAT 105 & PHYS 110)")

    ws.merge_cells(start_row=3, start_column=12, end_row=3, end_column=17)
    ws.cell(row=3, column=12, value="Applied Sociology & Applied Psychology (SOCI 115 & PSYC 120)")

    ws.merge_cells(start_row=3, start_column=18, end_row=3, end_column=20)
    ws.cell(row=3, column=18, value="Communicative English (ENGL 101)")

    # Row 4: Sub-component Headers (Demographic & Sub-component Anchors)
    headers = [
        (1, "S.No."),
        (2, "Student Name"),
        (3, "Father's Name"),
        (4, "Enroll No."),
        (5, "Roll NO."),
        # ANAT_PHYS
        (6, "Int. 25"),
        (7, "Anat Ex. 37"),
        (8, "Phys Ex. 38"),
        (9, "Ex. Total 75"),
        (10, "Grace"),
        (11, "Course Total 100"),
        # SOCI_PSYC
        (12, "Int. 25"),
        (13, "Soci Ex. 37"),
        (14, "Psyc Ex. 38"),
        (15, "Ex. Total 75"),
        (16, "Grace"),
        (17, "Course Total 100"),
        # ENGL
        (18, "Int. 25"),
        (19, "Ex. 25"),
        (20, "Total 50"),
        (21, "Result / Remarks"),
    ]
    for c, label in headers:
        ws.cell(row=4, column=c, value=label)

    # Row 5: Candidate with raw grace expression "35+2.5-37.5"
    row5 = [
        1, "ADITYA KUMAR KATHERIYA", "VIJAY KUMAR", 2025630, 202501002,
        20, 19, 20, 39, 0, 59,
        20, 20, 15, "35+2.5-37.5", 0, 57.5,
        21, 23, 44,
        "PASS WITH 2.5 GRACE MARKS IN APPLIED SOCIOLOGY & APPLIED PSYCHOLOGY"
    ]
    for c, val in enumerate(row5, 1):
        ws.cell(row=5, column=c, value=val)

    # Row 6: Candidate with "37+0.5=37.5"
    row6 = [
        2, "AMAN CHAUDHARY", "SURESH CHANDRA", "2025633.0", "202501005.0",
        15, 21, 20, 41, 0, 56,
        15, 24, 13, "37+0.5=37.5", 0, 52.5,
        20, 18, 38,
        "PASS WITH 0.5 GRACE MARKS IN APPLIED SOCIOLOGY & APPLIED PSYCHOLOGY"
    ]
    for c, val in enumerate(row6, 1):
        ws.cell(row=6, column=c, value=val)

    # Row 7: Footer signature row (Should trigger truncation!)
    ws.cell(row=7, column=1, value="Prepared By")
    ws.cell(row=7, column=5, value="Checked by: Mr Vinod Kumar Gupta")
    ws.cell(row=7, column=10, value="Dr. Radha K. (Sub-Dean)")

    # Row 8: Stray row that should NOT be parsed
    ws.cell(row=8, column=2, value="GHOST STUDENT THAT MUST NOT APPEAR")

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    adapter = ExamCellParserAdapter(file_bytes)
    records = adapter.parse()

    # Must only parse rows 5 and 6 (Row 7 triggers footer truncation!)
    assert len(records) == 2

    # Check Candidate 1 (ADITYA KUMAR KATHERIYA)
    r1 = records[0]
    assert r1.roll_no == "202501002"
    assert r1.enroll_no == "2025630"
    assert r1.student_name == "ADITYA KUMAR KATHERIYA"
    assert r1.father_name == "VIJAY KUMAR"
    assert "ANAT_PHYS" in r1.courses
    assert "SOCI_PSYC" in r1.courses
    assert "ENGL" in r1.courses

    # ANAT_PHYS marks
    anat = r1.courses["ANAT_PHYS"]
    assert anat["int_marks"] == 20
    assert anat["ex_part_a"] == 19
    assert anat["ex_part_b"] == 20
    assert anat["ex_total"] == 39
    assert anat["grace"] == 0
    assert anat["total"] == 59

    # SOCI_PSYC marks (extracted from "35+2.5-37.5")
    soci = r1.courses["SOCI_PSYC"]
    assert soci["int_marks"] == 20
    assert soci["ex_part_a"] == 20
    assert soci["ex_part_b"] == 15
    assert soci["grace"] == 2.5
    assert soci["total"] == 57.5

    # ENGL marks
    engl = r1.courses["ENGL"]
    assert engl["int_marks"] == 21
    assert engl["ex_marks"] == 23
    assert engl["total"] == 44

    assert "2.5 GRACE" in r1.raw_remarks

    # Check Candidate 2 (AMAN CHAUDHARY with float cleaned strings)
    r2 = records[1]
    assert r2.roll_no == "202501005"
    assert r2.enroll_no == "2025633"
    assert r2.courses["SOCI_PSYC"]["grace"] == 0.5


def test_parser_against_master_datasheet():
    """Tests adapter against the actual VidyaMarksheet_Master_DataSheet.xlsx file."""
    import pytest
    from pathlib import Path
    candidate_paths = [
        Path("VidyaMarksheet_Master_DataSheet.xlsx"),
        Path(__file__).resolve().parent.parent / "VidyaMarksheet_Master_DataSheet.xlsx",
        Path(__file__).resolve().parent.parent.parent / "VidyaMarksheet_Master_DataSheet.xlsx",
    ]
    path = next((p for p in candidate_paths if p.exists()), None)
    if not path:
        pytest.skip("VidyaMarksheet_Master_DataSheet.xlsx not found")
    with open(path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read(), sheet_name="Semester_1")
        records = adapter.parse()

    assert len(records) == 3
    aditya = next(r for r in records if r.roll_no == "202501002")
    assert aditya.student_name == "ADITYA KUMAR KATHERIYA"
    assert aditya.father_name == "VIJAY KUMAR"
    assert aditya.enroll_no == "2025630"
    assert aditya.courses["ANAT_PHYS"]["total"] == 59
    assert aditya.courses["SOCI_PSYC"]["grace"] == 2.5
    assert aditya.courses["SOCI_PSYC"]["total"] == 57.5
    assert aditya.courses["ENGL"]["total"] == 44
    assert "2.5 GRACE" in aditya.raw_remarks


def test_parser_against_vinod_sir_master_data():
    """Tests adapter and evaluation engine against the official master file from Mr. Vinod Sir."""
    import os
    from backend.evaluation_engine import EvaluationEngine

    file_path = r"E:\Examination\MARKS DATA\Received From Mr. Vinod Sir\First Year First Sem-JAN2026\1st Year Semester Master Data.xlsx"
    if not os.path.exists(file_path):
        pytest.skip("Vinod Sir master data file not found at expected path")

    with open(file_path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read(), sheet_name="Semester -I")
        records = adapter.parse()

    # Must parse all 67 candidates across all 5 pages (skipping intermediate signature blocks)
    assert len(records) == 67

    # Akash Verma (Roll 202501004) must pass with 185 marks and SGPA 7.00
    akash = next(r for r in records if r.roll_no == "202501004")
    assert akash.student_name == "AKASH VERMA"
    assert akash.father_name == "MAHESH CHANDRA"
    assert akash.enroll_no == "2025632"

    anat = akash.courses["ANAT_PHYS"]
    assert anat["int_marks"] == 19
    assert anat["ex_part_a"] == 23
    assert anat["ex_part_b"] == 22
    assert anat["ex_total"] == 45
    assert anat["total"] == 64

    soci = akash.courses["SOCI_PSYC"]
    assert soci["int_marks"] == 18
    assert soci["ex_part_a"] == 32
    assert soci["ex_part_b"] == 26
    assert soci["ex_total"] == 58
    assert soci["total"] == 76

    engl = akash.courses["ENGL"]
    assert engl["int_marks"] == 22
    assert engl["ex_marks"] == 23
    assert engl["total"] == 45

    engine = EvaluationEngine(1)
    eval_res = engine.evaluate_student(akash)
    assert eval_res.result_status == "PASS"
    assert eval_res.sgpa == 7.00
    assert eval_res.total_marks_obtained == 185.0
    assert eval_res.total_marks_max == 250.0
    assert len(eval_res.failed_courses) == 0


def test_unmerge_and_replicate_cells():
    """Verifies that unmerge_and_replicate_cells populates all merged coordinates."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("B2:D3")
    ws.cell(2, 2, value="Parent Title")

    unmerge_and_replicate_cells(ws)

    assert len(ws.merged_cells.ranges) == 0
    for r in range(2, 4):
        for c in range(2, 5):
            assert ws.cell(r, c).value == "Parent Title"


def test_parser_against_vinod_sir_sem2_master_data():
    """
    Tests deterministic merged-cell parser and evaluation engine against the official
    2nd Semester Master Data spreadsheet from Mr. Vinod Gupta Sir.
    Verifies all 64 candidates pass with 0 parsing errors.
    """
    import os
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    candidate_paths = [
        r"G:\Examination\MARKS DATA\21. All DataSheet Received From Mr. Vinod Gupta Sir on 7-10-26\2nd Semester Main & Re-appear Master Data July, 2026.xlsx",
        r"E:\Examination\MARKS DATA\21. All DataSheet Received From Mr. Vinod Gupta Sir on 7-10-26\2nd Semester Main & Re-appear Master Data July, 2026.xlsx",
    ]
    file_path = next((p for p in candidate_paths if os.path.exists(p)), None)
    if not file_path:
        pytest.skip("2nd Semester master data file not found at expected path")

    with open(file_path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read())
        records = adapter.parse()

    # Must parse all 64 candidates across all pages
    assert len(records) == 64

    # Candidate 1: Aditi Sharma (Roll 202502001)
    aditi = next(r for r in records if r.roll_no == "202502001")
    assert aditi.student_name == "ADITI SHARMA"
    assert aditi.enroll_no == "2025629"
    assert "BIOC_NUTR" in aditi.courses
    assert "NNF_THEORY" in aditi.courses
    assert "NNF_PRACT" in aditi.courses
    assert "HNIT" in aditi.courses

    # Candidate 58: Swati Shukla (Roll 202502059) with 2.5 grace
    swati = next(r for r in records if r.roll_no == "202502059")
    assert swati.student_name == "SWATI SHUKLA"
    assert swati.courses["BIOC_NUTR"]["grace"] == 2.5
    assert "2.5 GRACE" in swati.raw_remarks.upper()

    # Full Evaluation Engine Verification
    engine = EvaluationEngine(2)
    eval_res = engine.evaluate_batch(records)
    assert eval_res.total_candidates == 64
    assert eval_res.passed_count == 64
    assert eval_res.failed_count == 0
    assert eval_res.grace_awarded_count == 1

    swati_eval = next(r for r in eval_res.results if r.roll_no == "202502059")
    assert swati_eval.sgpa == 5.81
    assert "PASS WITH 2.5 GRACE MARKS" in swati_eval.result_status

    aditi_eval = next(r for r in eval_res.results if r.roll_no == "202502001")
    assert aditi_eval.sgpa == 7.55
    assert aditi_eval.result_status == "PASS"

    # Quality Audit Gate Verification
    gate = QualityAuditGate(semester_number=2)
    audit_report = gate.audit_batch(records, eval_res.results)
    assert audit_report.is_valid is True
    assert audit_report.can_export_pdf is True
    assert audit_report.critical_errors_count == 0
    assert audit_report.flagged_candidates == 0


def test_parser_against_vinod_sir_sem7_master_data():
    """
    Tests merged-cell parser, evaluation engine, and audit gate against official
    7th Semester Main Master Data spreadsheet (Feb 2026).
    Verifies 44 candidates: 39 Passed, 5 Failed, 1 Grace Awarded.
    """
    import os
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    candidate_paths = [
        r"G:\Examination\MARKS DATA\16.Data sheet B.Sc.Nursing Fourth Year (Seventh Semester) Examination held in Febuary-2026\7th Semester Master Data.xlsx",
        r"E:\Examination\MARKS DATA\16.Data sheet B.Sc.Nursing Fourth Year (Seventh Semester) Examination held in Febuary-2026\7th Semester Master Data.xlsx",
    ]
    file_path = next((p for p in candidate_paths if os.path.exists(p)), None)
    if not file_path:
        pytest.skip("7th Semester master data file not found at expected path")

    with open(file_path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read())
        records = adapter.parse()

    assert len(records) == 44

    # Pratham Verma (Roll 202207031) has 4.5 grace in OBGN_THEORY
    pratham = next(r for r in records if r.roll_no == "202207031")
    assert pratham.student_name == "PRATHAM VERMA"
    assert pratham.courses["OBGN_THEORY"]["grace"] == 4.5
    assert pratham.courses["OBGN_THEORY"]["ex_marks"] == 33

    # Evaluation Engine Verification
    engine = EvaluationEngine(7)
    eval_res = engine.evaluate_batch(records)
    assert eval_res.total_candidates == 44
    assert eval_res.passed_count == 39
    assert eval_res.failed_count == 5
    assert eval_res.grace_awarded_count == 1

    pratham_eval = next(r for r in eval_res.results if r.roll_no == "202207031")
    assert "PASS WITH 4.5 GRACE MARKS" in pratham_eval.result_status

    failed_rolls = {r.roll_no for r in eval_res.results if r.result_status == "FAIL"}
    assert failed_rolls == {"202207005", "202207011", "202207019", "202207026", "202207032"}

    # Quality Audit Gate Verification
    gate = QualityAuditGate(semester_number=7)
    audit_report = gate.audit_batch(records, eval_res.results)
    assert audit_report.is_valid is True
    assert audit_report.can_export_pdf is True
    assert audit_report.critical_errors_count == 0


def test_parser_against_vinod_sir_sem7_reappear_master_data():
    """
    Tests merged-cell parser and evaluation engine against official
    7th Semester Main & Re-appear Master Data spreadsheet (July 2026).
    Verifies 6 candidates (1 Main + 5 Re-appear): all 6 cleared, 1 Grace Awarded.
    """
    import os
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    candidate_paths = [
        r"G:\Examination\MARKS DATA\Data sheet B.Sc. Nursing 4th year(7th semester) (Main & Re-appear) July 2026\7th Semester Master Data re-appear.xlsx",
        r"E:\Examination\MARKS DATA\Data sheet B.Sc. Nursing 4th year(7th semester) (Main & Re-appear) July 2026\7th Semester Master Data re-appear.xlsx",
    ]
    file_path = next((p for p in candidate_paths if os.path.exists(p)), None)
    if not file_path:
        pytest.skip("7th Semester re-appear master data file not found at expected path")

    with open(file_path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read())
        records = adapter.parse()

    assert len(records) == 6

    # Deepak Yadav (Roll 202207016) with 1.5 grace in Nursing Research & Statistics
    deepak = next(r for r in records if r.roll_no == "202207016")
    assert deepak.student_name == "DEEPAK YADAV"

    # Evaluation Engine Verification
    engine = EvaluationEngine(7)
    eval_res = engine.evaluate_batch(records)
    assert eval_res.total_candidates == 6
    assert eval_res.passed_count == 6
    assert eval_res.failed_count == 0
    assert eval_res.grace_awarded_count == 1

    deepak_eval = next(r for r in eval_res.results if r.roll_no == "202207016")
    assert "PASS WITH 1.5 GRACE MARKS" in deepak_eval.result_status
    assert deepak_eval.total_grace == 1.5

    # Re-appear students all cleared
    reappear_rolls = ["202207005-S1", "202207011-S1", "202207019-S1", "202207026-S1", "202207032-S1"]
    for roll in reappear_rolls:
        r_eval = next(r for r in eval_res.results if r.roll_no == roll)
        assert r_eval.result_status == "PASS"

    # Quality Audit Gate Verification
    gate = QualityAuditGate(semester_number=7)
    audit_report = gate.audit_batch(records, eval_res.results)
    assert audit_report.is_valid is True
    assert audit_report.can_export_pdf is True
    assert audit_report.critical_errors_count == 0


def test_parser_against_vinod_sir_sem8_master_data():
    """
    Tests merged-cell parser and evaluation engine against official
    8th Semester Master Data spreadsheet (Competency Assessment).
    Verifies all 45 candidates pass with 0 errors.
    """
    import os
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    candidate_paths = [
        r"G:\Examination\MARKS DATA\20. B.Sc.Nursing Fourth Year (Eighth Semester) 2022-2026 Batch\4. Fourth Year\8th Semester\8th Semester Master Data.xlsx",
        r"E:\Examination\MARKS DATA\20. B.Sc.Nursing Fourth Year (Eighth Semester) 2022-2026 Batch\4. Fourth Year\8th Semester\8th Semester Master Data.xlsx",
    ]
    file_path = next((p for p in candidate_paths if os.path.exists(p)), None)
    if not file_path:
        pytest.skip("8th Semester master data file not found at expected path")

    with open(file_path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read())
        records = adapter.parse()

    assert len(records) == 45

    # Candidate 1: Abhinandan Kumar Sharma (Roll 202208001)
    abhinandan = next(r for r in records if r.roll_no == "202208001")
    assert abhinandan.student_name == "ABHINANDAN KUMAR SHARMA"
    assert "COMPETENCY" in abhinandan.courses
    assert abhinandan.courses["COMPETENCY"]["int_marks"] == 94
    assert abhinandan.courses["COMPETENCY"]["ex_marks"] == 95
    assert abhinandan.courses["COMPETENCY"]["total"] == 189

    # Evaluation Engine Verification
    engine = EvaluationEngine(8)
    eval_res = engine.evaluate_batch(records)
    assert eval_res.total_candidates == 45
    assert eval_res.passed_count == 45
    assert eval_res.failed_count == 0
    assert eval_res.grace_awarded_count == 0

    abhinandan_eval = next(r for r in eval_res.results if r.roll_no == "202208001")
    assert abhinandan_eval.result_status == "PASS"
    assert abhinandan_eval.sgpa == 10.0

    # Quality Audit Gate Verification
    gate = QualityAuditGate(semester_number=8)
    audit_report = gate.audit_batch(records, eval_res.results)
    assert audit_report.is_valid is True
    assert audit_report.can_export_pdf is True
    assert audit_report.critical_errors_count == 0


def test_parser_against_vinod_sir_sem4_master_data_and_3section_split():
    """
    Tests Semester 4 3-section split (PHAR 38 + PATH 25 + GEN 12) & Quality Audit Gate.
    Verifies:
    - 33 (PHAR) + 21 (PATH) + 8 (GEN) = 62 (Ext Total)
    - 20 (Int) + 62 (Ext) = 82 (Total Marks, Grade A+, GP 9)
    - Grand Total = 286 / 350
    - Quality Audit Gate passes with 0 critical errors (all 67 false criticals cleared).
    """
    import os
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    candidate_paths = [
        r"G:\Examination\MARKS DATA\21. All DataSheet Received From Mr. Vinod Gupta Sir on 7-10-26\4th Semester Master Data, July 2026.xlsx",
        r"E:\Examination\MARKS DATA\21. All DataSheet Received From Mr. Vinod Gupta Sir on 7-10-26\4th Semester Master Data, July 2026.xlsx",
    ]
    file_path = next((p for p in candidate_paths if os.path.exists(p)), None)
    if not file_path:
        pytest.skip("4th Semester master data file not found at expected path")

    with open(file_path, "rb") as f:
        adapter = ExamCellParserAdapter(f.read(), semester=4)
        records = adapter.parse()

    assert len(records) == 65

    # 1. Candidate 1: Aaraf-Ul-Islam (Roll 202404001)
    aaraf = next(r for r in records if r.roll_no == "202404001")
    assert aaraf.student_name == "Aaraf-Ul-Islam"
    phar_path = aaraf.courses["PHAR_PATH"]
    assert phar_path["ex_part_a"] == 33  # PHAR Theory
    assert phar_path["ex_part_b"] == 21  # PATH Theory
    assert phar_path["ex_part_c"] == 8   # GEN Theory
    assert phar_path["ex_total"] == 62   # 33 + 21 + 8 = 62
    assert phar_path["int_marks"] == 20
    assert phar_path["total"] == 82      # 20 + 62 = 82

    # 2. Evaluation Engine Verification
    engine = EvaluationEngine(4)
    eval_res = engine.evaluate_batch(records)
    assert eval_res.total_candidates == 65
    assert eval_res.passed_count == 63
    assert eval_res.failed_count == 2  # Amisha Vishwakarma & Shubham Kumar

    aaraf_eval = next(r for r in eval_res.results if r.roll_no == "202404001")
    assert aaraf_eval.result_status == "PASS"
    assert aaraf_eval.total_marks_obtained == 286.0
    assert aaraf_eval.total_marks_max == 350.0
    c_eval = aaraf_eval.courses["PHAR_PATH"]
    assert c_eval.letter_grade == "A+"
    assert c_eval.grade_point == 9

    # 3. Quality Audit Gate Verification (0 critical errors, all 67 cleared)
    gate = QualityAuditGate(semester_number=4)
    audit_report = gate.audit_batch(records, eval_res.results)
    assert audit_report.is_valid is True
    assert audit_report.can_export_pdf is True
    assert audit_report.critical_errors_count == 0
    assert len(audit_report.anomalies) == 0


def test_sem3_ahn_i_theory_column_mapping_and_parity_lock_resolution():
    """
    Verifies that for Semester 3 Adult Health Nursing - I (N-AHN (I) 215):
    - Col 13 'Int. (25)' binds to internal_obtained (17.0)
    - Col 14 'Total (75)' binds to external_obtained (49.0)
    - Col 15 'Total (100)' binds to total_obtained (66.0)
    - Total (75) is NEVER mistakenly mapped to total_obtained.
    - Quality Audit Gate verifies IA (17.0) + EA (49.0) = 66.0 == Sheet Total (66.0),
      resulting in 0 Critical Parity Violations and unlocked marksheet generation.
    """
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Semester -I"

    # Row 1: Blank
    # Row 2: Course Headers
    ws.cell(row=2, column=1, value="S.No.")
    ws.cell(row=2, column=2, value="Name")
    ws.cell(row=2, column=3, value="Father's Name")
    ws.cell(row=2, column=4, value="Enroll No.")
    ws.cell(row=2, column=5, value="Roll NO.")
    ws.cell(row=2, column=6, value="Applied Microbiology & Infection Control including patient safety (Code No. MICR-201)")
    ws.cell(row=2, column=13, value="Adult Health Nursing - I (Code No. N-AHN (I) 215)")
    ws.cell(row=2, column=18, value="Adult Health Nursing - I (Practical) (Code No. N-AHN (I) 215)")
    ws.cell(row=2, column=23, value="Grand Total (300)")
    ws.cell(row=2, column=24, value="Result/ Remark")

    # Row 3: Subcomponent Headers
    ws.cell(row=3, column=6, value="Int. 25")
    ws.cell(row=3, column=7, value="MICRO (Ex. 37)")
    ws.cell(row=3, column=8, value="INFC (Ex. 38)")
    ws.cell(row=3, column=9, value="Total (75)")
    ws.cell(row=3, column=10, value="Total (100)")
    ws.cell(row=3, column=11, value="Letter Grade")
    ws.cell(row=3, column=12, value="Grade Point")

    # The exact columns from the bug report
    ws.cell(row=3, column=13, value="Int. (25)")
    ws.cell(row=3, column=14, value="Total (75)")
    ws.cell(row=3, column=15, value="Total (100)")
    ws.cell(row=3, column=16, value="Letter Grade")
    ws.cell(row=3, column=17, value="Grade Point")

    ws.cell(row=3, column=18, value="Int. (50)")
    ws.cell(row=3, column=19, value="Ex. (50)")
    ws.cell(row=3, column=20, value="Total (100)")
    ws.cell(row=3, column=21, value="Letter Grade")
    ws.cell(row=3, column=22, value="Grade Point")

    # Row 4: Student Data (Abhinandan Kumar Sharma)
    ws.cell(row=4, column=1, value=1)
    ws.cell(row=4, column=2, value="ABHINANDAN KUMAR SHARMA")
    ws.cell(row=4, column=3, value="KANHAIYA LAL SHARMA")
    ws.cell(row=4, column=4, value=2022451)
    ws.cell(row=4, column=5, value=202203001)

    # MICR
    ws.cell(row=4, column=6, value=17)
    ws.cell(row=4, column=7, value=25)
    ws.cell(row=4, column=8, value=26)
    ws.cell(row=4, column=9, value=51)
    ws.cell(row=4, column=10, value=68)
    ws.cell(row=4, column=11, value="B+")
    ws.cell(row=4, column=12, value=7)

    # AHN_I_THEORY
    ws.cell(row=4, column=13, value=17.0)
    ws.cell(row=4, column=14, value=49.0)
    ws.cell(row=4, column=15, value=66.0)
    ws.cell(row=4, column=16, value="B+")
    ws.cell(row=4, column=17, value=7)

    # AHN_I_PRACT
    ws.cell(row=4, column=18, value=37.0)
    ws.cell(row=4, column=19, value=35.0)
    ws.cell(row=4, column=20, value=72.0)
    ws.cell(row=4, column=21, value="B+")
    ws.cell(row=4, column=22, value=7)

    ws.cell(row=4, column=23, value=206.0)
    ws.cell(row=4, column=24, value="Pass")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    adapter = ExamCellParserAdapter(buf.getvalue(), semester=3)
    records = adapter.parse()

    assert len(records) == 1
    rec = records[0]
    assert rec.student_name == "ABHINANDAN KUMAR SHARMA"

    # Assert AHN Theory mapping
    ahn_theory = rec.courses["AHN_I_THEORY"]
    assert ahn_theory["int_marks"] == 17.0, f"Expected 17.0, got {ahn_theory['int_marks']}"
    assert ahn_theory["ex_marks"] == 49.0, f"Expected 49.0, got {ahn_theory['ex_marks']}"
    assert ahn_theory["total"] == 66.0, f"Expected 66.0, got {ahn_theory['total']} (Bug: mapped Total (75) to total!)"

    # Evaluation & Audit
    engine = EvaluationEngine(3)
    eval_res = engine.evaluate_batch(records)
    assert eval_res.results[0].result_status == "PASS"

    gate = QualityAuditGate(3)
    audit = gate.audit_batch(records, eval_res.results)
    assert audit.is_valid is True
    assert audit.critical_errors_count == 0
    assert audit.can_export_pdf is True


def test_sem3_real_master_data_ingestion_and_zero_parity_errors():
    """
    Tests against real 3rd Semester Master Data files on disk if present.
    Verifies 0 critical parity errors across all candidates.
    """
    import os
    from backend.evaluation_engine import EvaluationEngine
    from backend.audit_engine import QualityAuditGate

    candidate_paths = [
        r"E:\Examination\MARKS DATA\20. B.Sc.Nursing Fourth Year (Eighth Semester) 2022-2026 Batch\2. Second Year\3rd Semester\3rd Semester Master Data.xlsx",
        r"E:\Examination\MARKS DATA\5.Data sheet B.Sc. Nursing Second Year Third Semester (Main & Re-appear) Examination Held in the Month of January 2026\3rd Semester Master Data.xlsx",
    ]
    existing_paths = [p for p in candidate_paths if os.path.exists(p)]
    if not existing_paths:
        pytest.skip("No real 3rd Semester Master Data file found on disk.")

    for path in existing_paths:
        with open(path, "rb") as f:
            adapter = ExamCellParserAdapter(f.read(), semester=3)
            records = adapter.parse()

        assert len(records) > 0
        engine = EvaluationEngine(3)
        eval_res = engine.evaluate_batch(records)
        gate = QualityAuditGate(3)
        audit = gate.audit_batch(records, eval_res.results)

        assert audit.is_valid is True, f"Parity errors found in {path}: {audit.anomalies}"
        assert audit.critical_errors_count == 0
        assert audit.can_export_pdf is True



