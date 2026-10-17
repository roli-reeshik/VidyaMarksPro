"""
VidyaMarksheet Backlog & Re-appear Management Route
Compliant with INC Gazette Corrigendum & SGPGIMS Lucknow Regulations.
Branding: VidyaMarksheet (Developed at VidyaLabs | Priyanka Kumari)

Endpoints:
- GET /api/v1/backlogs/reappear-roster
- GET /api/v1/backlogs/export-roster-excel
"""

import io
import logging
from datetime import datetime
from typing import Optional, List, Dict, Any

from fastapi import APIRouter, Depends, Query, HTTPException, Response, status
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.db.database import get_sync_session
from backend.db.repository import ExaminationRepository
from backend.security import require_authenticated

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1/backlogs", tags=["Backlogs & Re-appear"])


@router.get("/reappear-roster")
def get_reappear_roster(
    semester: Optional[int] = Query(None, ge=1, le=8, description="Target Semester (1-8)"),
    forthcoming_session: Optional[str] = Query(None, description="e.g. 'JULY - 2026', 'JANUARY - 2027'"),
    status: Optional[str] = Query("ACTIVE", description="'ACTIVE', 'CLEARED', or 'ALL'"),
    current_user: Any = Depends(require_authenticated),
):
    """
    Retrieves the authoritative roster of students with active or historical backlogs,
    including deficit marks, failed component, attempt count, and forthcoming examination sitting tags.
    """
    with get_sync_session() as session:
        repo = ExaminationRepository(session)
        # Ensure any historical unregistered failed marks are synchronized
        repo.sync_all_unregistered_backlogs()
        roster = repo.get_backlog_roster(
            semester_number=semester,
            backlog_status=status or "ACTIVE",
            forthcoming_session=forthcoming_session,
        )
        return {
            "status": "success",
            "total_candidates": len(roster),
            "semester_filter": semester,
            "forthcoming_session_filter": forthcoming_session,
            "status_filter": status,
            "roster": roster,
            "branding": "VidyaMarksheet (Developed at VidyaLabs | Priyanka Kumari)",
        }


@router.get("/export-roster-excel")
def export_reappear_roster_excel(
    semester: Optional[int] = Query(None, ge=1, le=8, description="Target Semester (1-8)"),
    forthcoming_session: Optional[str] = Query(None, description="e.g. 'JULY - 2026', 'JANUARY - 2027'"),
    status: Optional[str] = Query("ACTIVE", description="'ACTIVE', 'CLEARED', or 'ALL'"),
    current_user: Any = Depends(require_authenticated),
):
    """
    Generates and downloads the official administrative Re-appear Candidates Roster
    formatted in Microsoft Excel (.xlsx) with institutional header, styling, and column metrics.
    """
    try:
        with get_sync_session() as session:
            repo = ExaminationRepository(session)
            repo.sync_all_unregistered_backlogs()
            roster = repo.get_backlog_roster(
                semester_number=semester,
                backlog_status=status or "ACTIVE",
                forthcoming_session=forthcoming_session,
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Re-appear Roster"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        sub_navy_fill = PatternFill(start_color="2B4C7E", end_color="2B4C7E", fill_type="solid")
        zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        header_fill = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")

        title_font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
        sub_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        meta_font = Font(name="Arial", size=9, italic=True, color="475569")
        col_hdr_font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        data_font = Font(name="Arial", size=9, bold=False, color="1E293B")
        bold_data_font = Font(name="Arial", size=9, bold=True, color="1E293B")
        fail_font = Font(name="Arial", size=9, bold=True, color="DC2626")

        thin_side = Side(border_style="thin", color="CBD5E1")
        cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)

        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center")

        # Row 1: Institutional Header
        ws.merge_cells("A1:O1")
        ws["A1"] = "SANJAY GANDHI POSTGRADUATE INSTITUTE OF MEDICAL SCIENCES, LUCKNOW"
        ws["A1"].font = title_font
        ws["A1"].fill = navy_fill
        ws["A1"].alignment = center_align
        ws.row_dimensions[1].height = 28

        # Row 2: Subtitle
        ws.merge_cells("A2:O2")
        ws["A2"] = "COLLEGE OF NURSING — OFFICIAL RE-APPEAR CANDIDATES ROSTER & BACKLOG SCHEDULE"
        ws["A2"].font = sub_font
        ws["A2"].fill = sub_navy_fill
        ws["A2"].alignment = center_align
        ws.row_dimensions[2].height = 20

        # Row 3: Metadata Block
        ws.merge_cells("A3:O3")
        filter_str = (
            f"Semester: {semester if semester else 'All Semesters'} | "
            f"Forthcoming Session: {forthcoming_session if forthcoming_session else 'All Forthcoming Sessions'} | "
            f"Status Filter: {status or 'ACTIVE'} | "
            f"Total Candidates: {len(roster)} | "
            f"Generated: {datetime.now().strftime('%d-%b-%Y %H:%M IST')} | "
            f"VidyaMarksheet (Developed at VidyaLabs | Priyanka Kumari)"
        )
        ws["A3"] = filter_str
        ws["A3"].font = meta_font
        ws["A3"].alignment = left_align
        ws.row_dimensions[3].height = 18

        # Row 4: Column Headers
        headers = [
            "S.No.",
            "Roll No.",
            "Enrollment No.",
            "Student Name",
            "Father's Name",
            "Failed Sem",
            "Subject Code",
            "Subject Title",
            "Failed Component",
            "Deficit Marks",
            "Attempt",
            "Re-appear Sitting Tag",
            "Failed Exam Session",
            "Forthcoming Exam Session",
            "Backlog Status",
        ]

        ws.row_dimensions[4].height = 24
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h_text)
            cell.font = col_hdr_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = cell_border

        # Data Rows
        current_row = 5
        for idx, item in enumerate(roster, start=1):
            ws.row_dimensions[current_row].height = 20
            row_fill = zebra_fill if idx % 2 == 0 else white_fill

            row_values = [
                idx,
                item.get("roll_no", ""),
                item.get("enrollment_no", ""),
                item.get("student_name", ""),
                item.get("father_name", ""),
                f"Sem {item.get('failed_semester', '')}",
                item.get("failed_subject_code", ""),
                item.get("failed_subject_title", ""),
                item.get("component_failed", ""),
                f"{item.get('deficit_marks', 0.0):.2f}",
                item.get("attempt_count", 1),
                item.get("attempt_label", "REAPPEAR - I"),
                item.get("failed_exam_month_year", ""),
                item.get("forthcoming_exam_month_year", ""),
                item.get("backlog_status", "ACTIVE"),
            ]

            for c_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.fill = row_fill
                cell.border = cell_border

                # Alignment & font rules
                if c_idx in [1, 2, 3, 6, 10, 11, 12, 13, 14, 15]:
                    cell.alignment = center_align
                elif c_idx in [4, 5, 8]:
                    cell.alignment = left_align
                elif c_idx == 7:
                    cell.alignment = left_align

                if c_idx in [2, 12]:
                    cell.font = bold_data_font
                elif c_idx == 10:
                    cell.font = fail_font
                elif c_idx == 15:
                    cell.font = fail_font if str(val).upper() == "ACTIVE" else bold_data_font
                else:
                    cell.font = data_font

            current_row += 1

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            # Skip checking the merged rows 1-3 for width
            for cell in col[3:]:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        # Output to buffer
        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        sem_suffix = f"Sem{semester}" if semester else "AllSemesters"
        session_suffix = (forthcoming_session or "AllSessions").replace(" ", "_").replace("-", "_")
        filename = f"SGPGIMS_Reappear_Roster_{sem_suffix}_{session_suffix}.xlsx"

        return Response(
            content=stream.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    except Exception as e:
        logger.error(f"Error generating Re-appear Roster Excel: {e}", exc_info=True)
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error generating Re-appear Roster Excel: {str(e)}",
        )
